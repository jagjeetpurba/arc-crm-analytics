"""The dashboard must follow the data: new column -> new widgets/filters; ERP change -> heal -> version bump."""
import openpyxl, shutil
from app.state import get_state
from tests.conftest import WORKBOOK


def test_new_column_appears_without_code_change(client, tmp_path):
    wb = openpyxl.load_workbook(WORKBOOK)            # normal (non read-only) so we can edit
    ws = wb["PILOT Opportunities"]
    hdr = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r[0] == "Opportunity GUID")
    col = ws.max_column + 1
    ws.cell(hdr, col, "Sales Region")
    for i in range(hdr + 1, ws.max_row + 1):
        if ws.cell(i, 1).value:
            ws.cell(i, col, ["North", "South", "East"][i % 3])
    f = tmp_path / "v2.xlsx"; wb.save(f)
    v0 = get_state().version
    with open(f, "rb") as fh:
        r = client.post("/api/ingest/xlsx", files={"file": ("v2.xlsx", fh)})
    assert r.status_code == 200, r.text
    assert "Sales Region" in r.json()["report"]["opportunities"]["dynamic_headers"]
    assert get_state().version > v0
    spec = client.get("/api/spec").json()
    assert "sales_region" in spec["filters"]["dynamic"]
    assert any(w["id"] == "auto_opportunities_sales_region" for p in spec["pages"] for w in p["widgets"])
    d = client.get("/api/widget/auto_opportunities_sales_region").json()["payload"]
    assert {r["key"] for r in d["rows"]} == {"North", "South", "East"}
    # the new dimension filters the curated KPIs too
    a = client.get("/api/widget/exec_close_month", params={"period": "all", "sales_region": "North"}).json()["payload"]
    assert a["scope"]["filters_applied"].get("sales_region")


def test_erp_change_is_detected_and_healed(client):
    s = get_state()
    assert s.drift() == []
    row = next(r for r in s.data["opportunities"] if r["status_reason"] == "In Progress")
    s.erp.upsert("opportunities", [{**row, "est_gp": (row["est_gp"] or 0) + 1000}], source="erp-edit")   # someone edits the ERP
    assert s.drift() == ["opportunities"]
    st = client.get("/api/sync/status").json()
    assert st["drift"] == ["opportunities"]
    v = s.version
    r = client.post("/api/sync/heal").json()
    assert r["changed"] == ["opportunities"] and s.version == v + 1 and s.drift() == []


def test_predictive_exposes_drivers_and_limits(client):
    p = client.get("/api/widget/pr_drivers").json()["payload"]
    assert p["available"] and p["drivers"] and 0 <= p["model"]["cv_auc"] <= 1 and p["model"]["caveat"]
    assert p["forecast"]["p10"] <= p["forecast"]["p90"]
