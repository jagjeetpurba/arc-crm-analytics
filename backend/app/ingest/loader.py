"""Workbook -> ERP. Header-driven, idempotent, schema-drift tolerant.

* headers are matched to ERP fields via spec/erp_schema.yaml `source`
* headers that are neither mapped nor in ignore_headers are kept as dynamic `extra` fields
* re-running changes nothing (row hash per key)
"""
from __future__ import annotations
import re
from datetime import date, datetime
from pathlib import Path
import openpyxl
from ..erp.base import ERPClient


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def to_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return None


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat() if v.hour == v.minute == v.second == 0 else v.isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    return s.replace(" ", "T") if re.match(r"\d{4}-\d{2}-\d{2}", s) else None


def coerce(v, spec):
    t = spec["type"]
    if t == "float" or t == "int":
        return to_float(v)
    if t == "date":
        return to_date(v)
    if v is None:
        return None
    s = str(v).strip() if spec.get("trim") or isinstance(v, str) else str(v)
    return s or None


def find_header(ws, first_cell: str):
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row and norm(row[0]) == norm(first_cell) or (row and first_cell and norm(first_cell) in [norm(c) for c in row[:2]]):
            return i, list(row)
        if i > 15:
            break
    return None, None


def load_object(wb, name: str, obj: dict):
    sheet = obj.get("sheet")
    if not sheet or sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    hdr_i, headers = find_header(ws, obj["header_first_cell"])
    if not hdr_i:
        raise ValueError(f"[{name}] header row starting '{obj['header_first_cell']}' not found in '{sheet}'")
    by_source = {norm(c["source"]): (col, c) for col, c in obj["columns"].items() if "source" in c}
    ignore = {norm(h) for h in obj.get("ignore_headers", [])}
    plan = []   # (index, kind, col, spec)
    dynamic = []
    for idx, h in enumerate(headers):
        if h is None or str(h).strip() == "":
            continue
        n = norm(h)
        if n in by_source:
            plan.append((idx, "known", *by_source[n]))
        elif n in ignore or n.startswith("(do not modify)"):
            continue
        else:
            col = re.sub(r"[^a-z0-9]+", "_", n).strip("_")
            plan.append((idx, "dyn", col, {"type": "str"}))
            dynamic.append(str(h))
    missing = [c for c, s in obj["columns"].items() if "source" in s and norm(s["source"]) not in {norm(h) for h in headers if h}]
    rows, dyn_vals = [], {}
    for r in ws.iter_rows(min_row=hdr_i + 1, values_only=True):
        if not r or all(c is None or c == "" for c in r):
            continue
        row = {}
        for idx, kind, col, spec in plan:
            v = r[idx] if idx < len(r) else None
            if kind == "dyn":
                if isinstance(v, str) and v.startswith("="):
                    continue            # unevaluated formula: ignore
                dyn_vals.setdefault(col, []).append(v)
                row[col] = v if not isinstance(v, (datetime, date)) else to_date(v)
            else:
                row[col] = coerce(v, spec)
        if row.get(obj["key"]):
            rows.append(row)
    # numeric-ify dynamic columns whose values are all numeric
    for col, vals in dyn_vals.items():
        nonnull = [v for v in vals if v not in (None, "")]
        if nonnull and all(isinstance(v, (int, float)) or to_float(v) is not None and re.fullmatch(r"-?[\d,.]+", str(v)) for v in nonnull):
            for row in rows:
                if col in row:
                    row[col] = to_float(row[col])
    return rows, {"sheet": sheet, "dynamic_headers": dynamic, "missing_headers": missing}


def load_workbook_into_erp(path: Path, erp: ERPClient, schema: dict, dry_run: bool = False) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    report = {}
    for name, obj in schema["objects"].items():
        res = load_object(wb, name, obj)
        if res is None:
            continue
        rows, meta = res
        stats = {"inserted": 0, "updated": 0, "unchanged": 0} if dry_run else erp.upsert(name, rows, source=Path(path).name)
        report[name] = {"rows": len(rows), **stats, **meta}
    return report
